"""
파일 서버 스캔: 연결 조회 → 파일 다운로드 → 미디어 타입별 PII 탐지 → piiDetails 집계
"""
import base64
import tempfile
import os
from typing import List, Dict, Any
from pathlib import Path

from app.crud.file_server_connection import get_file_server_connection
from app.services.file.fetchers import fetch_file_bytes
from app.services.file.pii_slot_analyzer import (
    analyze_slots_for_masked,
    aggregate_document_pii,
    PII_TYPE_NAME_TO_CODE,
    PII_LABEL_TO_CODE,
)
from app.services.file.video_pii_dedup import dedup_video_faces, dedup_video_text_regions
from app.core.logging import logger

# 확장자 -> 미디어 타입 (document / image / audio / video)
DOC_EXT = {".txt", ".pdf", ".docx", ".doc", ".xlsx", ".xls"}
IMAGE_EXT = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp"}
AUDIO_EXT = {".wav", ".mp3", ".m4a", ".flac", ".ogg"}
VIDEO_EXT = {".mp4", ".avi", ".mov", ".mkv", ".webm"}


def _bytes_to_document_text(data: bytes, ext: str) -> str:
    """문서 바이트를 평문 텍스트로 변환"""
    ext_lower = ext.lower()
    if ext_lower == ".txt":
        return data.decode("utf-8", errors="replace")
    if ext_lower == ".pdf":
        try:
            import fitz
            doc = fitz.open(stream=data, filetype="pdf")
            text = "\n".join(page.get_text() for page in doc)
            doc.close()
            return text
        except Exception as e:
            logger.warning("PDF 텍스트 추출 실패: %s", e)
            return ""
    if ext_lower in (".docx", ".doc"):
        try:
            from docx import Document
            from io import BytesIO
            doc = Document(BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception as e:
            logger.warning("DOCX 텍스트 추출 실패: %s", e)
            return ""
    if ext_lower in (".xlsx", ".xls"):
        try:
            import openpyxl
            from io import BytesIO
            wb = openpyxl.load_workbook(BytesIO(data), read_only=True)
            parts = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    parts.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(parts)
        except Exception as e:
            logger.warning("엑셀 텍스트 추출 실패: %s", e)
            return ""
    return ""


def _aggregate_pii_details(
    document_detected: List[Dict] = None,
    image_face_count: int = 0,
    audio_detected: List[Dict] = None,
    video_result: Dict = None,
    document_masked: Dict[str, int] = None,
    video_face_count: int = None,
    video_text_pii_by_label: Dict[str, int] = None,
) -> List[Dict[str, Any]]:
    """탐지 결과를 piiDetails 형태로 집계 (piiType, totalCount, maskedCount). 영상은 트랙 단위 dedup 적용 시 video_face_count, video_text_pii_by_label 전달."""
    total_by_code: Dict[str, int] = {}
    masked_by_code: Dict[str, int] = dict(document_masked or {})

    if document_detected:
        for item in document_detected:
            t = item.get("type") or ""
            code = PII_TYPE_NAME_TO_CODE.get(t) or t
            total_by_code[code] = total_by_code.get(code, 0) + 1
    if image_face_count > 0:
        total_by_code["FACE"] = total_by_code.get("FACE", 0) + image_face_count
    if audio_detected:
        for item in audio_detected:
            t = item.get("type") or ""
            code = PII_TYPE_NAME_TO_CODE.get(t) or t
            total_by_code[code] = total_by_code.get(code, 0) + 1
    if video_result:
        faces = video_result.get("faces", [])
        audio_items = video_result.get("personal_info_in_audio", [])
        text_regions = video_result.get("text_pii_regions", [])
        if video_face_count is not None:
            if video_face_count > 0:
                total_by_code["FACE"] = total_by_code.get("FACE", 0) + video_face_count
        elif faces:
            total_by_code["FACE"] = total_by_code.get("FACE", 0) + len(faces)
        for item in audio_items:
            t = item.get("type") or ""
            code = PII_TYPE_NAME_TO_CODE.get(t) or t
            total_by_code[code] = total_by_code.get(code, 0) + 1
        if video_text_pii_by_label:
            for label, count in video_text_pii_by_label.items():
                if count <= 0:
                    continue
                code = PII_LABEL_TO_CODE.get(label) or PII_TYPE_NAME_TO_CODE.get(label) or label
                total_by_code[code] = total_by_code.get(code, 0) + count
        else:
            for r in text_regions:
                label = r.get("label", "p_nm")
                code = PII_LABEL_TO_CODE.get(label) or PII_TYPE_NAME_TO_CODE.get(label) or label
                total_by_code[code] = total_by_code.get(code, 0) + 1

    all_codes = set(total_by_code) | set(masked_by_code)
    return [
        {"piiType": code, "totalCount": total_by_code.get(code, 0), "maskedCount": masked_by_code.get(code, 0)}
        for code in sorted(all_codes)
    ]


def scan_files(
    connection_id: str,
    pii_files: List[str],
    document_detector,
    image_detector,
    audio_detector,
    video_detector,
) -> List[Dict[str, Any]]:
    """
    파일 서버에서 파일 목록을 받아 각 파일별 PII 탐지 결과 반환.

    Returns:
        [{"filePath": str, "piiDetected": bool, "piiDetails": [...]}, ...]
    """
    connection_info = get_file_server_connection(connection_id)
    results: List[Dict[str, Any]] = []

    for path in pii_files:
        file_path_display = path  # 요청 값 그대로 반환
        try:
            data = fetch_file_bytes(connection_info, path)
        except Exception as e:
            logger.warning(f"파일 다운로드 실패 path={path} error={e!r}", exc_info=True)
            results.append({
                "filePath": file_path_display,
                "piiDetected": False,
                "piiDetails": [],
            })
            continue

        ext = Path(path).suffix
        ext_lower = ext.lower()

        if ext_lower in DOC_EXT:
            text = _bytes_to_document_text(data, ext_lower)
            if not text.strip():
                results.append({"filePath": file_path_display, "piiDetected": False, "piiDetails": []})
                continue
            detected = document_detector.detect(text)
            masked_counts = analyze_slots_for_masked(text, document_detector)
            agg = aggregate_document_pii(detected, masked_counts)
            pii_details = [
                {"piiType": k, "totalCount": v["totalCount"], "maskedCount": v["maskedCount"]}
                for k, v in agg.items()
            ]
            pii_detected = any(d["totalCount"] > 0 or d["maskedCount"] > 0 for d in pii_details)
            results.append({"filePath": file_path_display, "piiDetected": pii_detected, "piiDetails": pii_details})
            continue

        if ext_lower in IMAGE_EXT:
            try:
                b64 = base64.b64encode(data).decode("ascii")
                faces = image_detector.detect_faces(b64)
            except Exception as e:
                logger.warning("이미지 탐지 실패 %s: %s", path, e)
                faces = []
            pii_details = [{"piiType": "FACE", "totalCount": len(faces), "maskedCount": 0}] if faces else []
            results.append({
                "filePath": file_path_display,
                "piiDetected": len(faces) > 0,
                "piiDetails": pii_details,
            })
            continue

        if ext_lower in AUDIO_EXT:
            try:
                b64 = base64.b64encode(data).decode("ascii")
                items = audio_detector.detect(b64, "base64")
            except Exception as e:
                logger.warning("오디오 탐지 실패 %s: %s", path, e)
                items = []
            pii_details = _aggregate_pii_details(audio_detected=items)
            results.append({
                "filePath": file_path_display,
                "piiDetected": len(items) > 0,
                "piiDetails": pii_details,
            })
            continue

        if ext_lower in VIDEO_EXT:
            try:
                with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                    tmp.write(data)
                    tmp_path = tmp.name
                try:
                    video_result = video_detector.detect(tmp_path, "path")
                finally:
                    try:
                        os.unlink(tmp_path)
                    except Exception:
                        pass
            except Exception as e:
                logger.warning("영상 탐지 실패 %s: %s", path, e)
                video_result = {}
            faces = video_result.get("faces", [])
            audio_items = video_result.get("personal_info_in_audio", [])
            text_regions = video_result.get("text_pii_regions", [])
            unique_face_count = dedup_video_faces(faces)
            text_pii_by_label = dedup_video_text_regions(text_regions)
            has_pii = unique_face_count > 0 or len(audio_items) > 0 or sum(text_pii_by_label.values()) > 0
            pii_details = _aggregate_pii_details(
                video_result=video_result,
                video_face_count=unique_face_count,
                video_text_pii_by_label=text_pii_by_label,
            )
            results.append({
                "filePath": file_path_display,
                "piiDetected": has_pii,
                "piiDetails": pii_details,
            })
            continue

        results.append({"filePath": file_path_display, "piiDetected": False, "piiDetails": []})

    return results
